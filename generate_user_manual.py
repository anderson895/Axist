"""
Generates a user manual (.docx) explaining how to change fonts, pictures,
text, and widths in the Axis website codebase.
"""

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def shade_cell(cell, color_hex):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    tc_pr.append(shd)


def add_heading(doc, text, level=1, color=(30, 60, 90)):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(*color)
        run.font.name = "Calibri"
    return h


def add_paragraph(doc, text, bold=False, size=11):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.bold = bold
    return p


def add_code_block(doc, code, language="tsx"):
    """Add a styled code block (monospace, light grey background)."""
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    cell = table.rows[0].cells[0]
    shade_cell(cell, "F4F4F4")
    cell.width = Inches(6.5)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(code)
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(40, 40, 40)
    doc.add_paragraph()


def add_screenshot_placeholder(doc, label):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"[ SCREENSHOT: {label} ]")
    run.italic = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(120, 120, 120)


def add_note(doc, text):
    p = doc.add_paragraph()
    run = p.add_run("NOTE: ")
    run.bold = True
    run.font.color.rgb = RGBColor(180, 90, 0)
    run2 = p.add_run(text)
    run2.font.size = Pt(11)


def add_step(doc, n, text):
    p = doc.add_paragraph(style="List Number")
    run = p.add_run(text)
    run.font.size = Pt(11)


def main():
    doc = Document()

    # Default style
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # ───────── Title ─────────
    title = doc.add_heading("Axis Website — Code Editing User Manual", level=0)
    for run in title.runs:
        run.font.color.rgb = RGBColor(20, 50, 80)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s = sub.add_run("How to change Fonts, Pictures, Text, and Width in the source code")
    s.italic = True
    s.font.size = Pt(12)
    s.font.color.rgb = RGBColor(90, 90, 90)

    doc.add_paragraph()

    # ───────── Table of contents ─────────
    add_heading(doc, "Contents", level=1)
    for i, item in enumerate([
        "1. Project Overview",
        "2. How to Change Fonts",
        "3. How to Change Pictures (Images)",
        "4. How to Change Text",
        "5. How to Change Width (Sizing & Layout)",
        "6. Quick Reference Cheat-Sheet",
    ], 1):
        p = doc.add_paragraph(item)
        p.paragraph_format.space_after = Pt(2)

    doc.add_page_break()

    # ───────── 1. Project Overview ─────────
    add_heading(doc, "1. Project Overview", level=1)
    add_paragraph(
        doc,
        "The Axis website is built with React + TypeScript + Vite, and uses "
        "Tailwind CSS for styling. The folder structure that you will edit most is:",
    )
    add_code_block(
        doc,
        "axis-website/\n"
        "├── index.html              ← Google Fonts imported here\n"
        "├── tailwind.config.js      ← Font families + colors registered here\n"
        "├── public/\n"
        "│   └── images/             ← All image files (pictures) live here\n"
        "│       └── roles/          ← Role card images\n"
        "└── src/\n"
        "    ├── index.css           ← Global styles & default body font\n"
        "    ├── data/\n"
        "    │   └── roles.ts        ← Text + image references for role cards\n"
        "    ├── pages/              ← Each page (LandingPage, RolesPage, etc.)\n"
        "    └── components/         ← Reusable UI parts (Navbar, Footer, …)\n",
    )
    add_screenshot_placeholder(doc, "VS Code Explorer panel showing folder structure")

    doc.add_page_break()

    # ───────── 2. Fonts ─────────
    add_heading(doc, "2. How to Change Fonts", level=1)
    add_paragraph(
        doc,
        "Fonts are configured in THREE places. To change the site's fonts you "
        "must update all three so they stay in sync.",
        bold=True,
    )

    add_heading(doc, "2.1  Import the font in index.html", level=2)
    add_paragraph(doc, "File: index.html (root of the project)")
    add_paragraph(doc, "Lines 8–11 import Google Fonts. Replace the URL on line 10 with the new font you want.")
    add_code_block(
        doc,
        '<link rel="preconnect" href="https://fonts.googleapis.com" />\n'
        '<link rel="preconnect" href="https://fonts.gstatic.com" crossorigin />\n'
        '<link href="https://fonts.googleapis.com/css2?family=Playfair+Display:'
        'ital,wght@0,400;0,700;0,900&family=DM+Sans:ital,wght@0,300;0,400;'
        '0,500;0,600;0,700&display=swap" rel="stylesheet" />\n'
        '<link href="https://fonts.googleapis.com/icon?family=Material+Icons" '
        'rel="stylesheet" />\n',
    )
    add_step(doc, 1, "Go to https://fonts.google.com and pick the new font.")
    add_step(doc, 2, "Click 'Get embed code' and copy the <link> tag they give you.")
    add_step(doc, 3, "Paste it in index.html replacing the existing <link> on line 10.")
    add_screenshot_placeholder(doc, "index.html opened in VS Code with line 10 highlighted")

    add_heading(doc, "2.2  Register the font in tailwind.config.js", level=2)
    add_paragraph(doc, "File: tailwind.config.js")
    add_paragraph(doc, "Lines 36–39 define the Tailwind font utilities. Update the values to match your new font.")
    add_code_block(
        doc,
        "fontFamily: {\n"
        "  display: ['\"Playfair Display\"', 'Georgia', 'serif'],   // → font-display\n"
        "  body:    ['\"DM Sans\"', '\"Segoe UI\"', 'sans-serif'],     // → font-body\n"
        "},\n",
    )
    add_note(
        doc,
        "Always wrap multi-word font names in escaped double-quotes "
        '(\\" \\") and keep at least one fallback like Georgia or sans-serif.',
    )
    add_screenshot_placeholder(doc, "tailwind.config.js with fontFamily block highlighted")

    add_heading(doc, "2.3  Update the default body font in src/index.css", level=2)
    add_paragraph(doc, "File: src/index.css")
    add_paragraph(doc, "Line 18 sets the base font that the whole page inherits.")
    add_code_block(
        doc,
        "body {\n"
        "  font-family: 'DM Sans', 'Segoe UI', sans-serif;\n"
        "  ...\n"
        "}\n",
    )
    add_screenshot_placeholder(doc, "src/index.css with the body font-family line highlighted")

    add_heading(doc, "2.4  Using the font in components", level=2)
    add_paragraph(doc, "Inside any .tsx file you can apply the registered fonts as Tailwind classes:")
    add_code_block(
        doc,
        '<h1 className="font-display text-4xl">Big Heading</h1>\n'
        '<p  className="font-body text-base">Regular paragraph text</p>\n',
    )

    doc.add_page_break()

    # ───────── 3. Pictures ─────────
    add_heading(doc, "3. How to Change Pictures (Images)", level=1)
    add_paragraph(
        doc,
        "All images live inside the public/ folder. Files placed there are "
        "served from the website root, so an image at "
        "public/images/roles/foo.png is reached using the URL /images/roles/foo.png.",
    )

    add_heading(doc, "3.1  Replace a role-card image", level=2)
    add_step(doc, 1, "Open the folder public/images/roles/ in your file explorer.")
    add_step(doc, 2, "Drop your new image inside (use .png or .jpg, ideally 800×600 or larger).")
    add_step(doc, 3, "Open src/data/roles.ts and find the role you want to update.")
    add_step(doc, 4, "Change the `image:` field to the new file's path.")
    add_code_block(
        doc,
        "{\n"
        "  title: 'Executive Assistant',\n"
        "  desc:  'Hire trusted, experienced remote executive assistants…',\n"
        "  savings: '86%',\n"
        "  image: '/images/roles/executive-assistant.png'   // ← change me\n"
        "},\n",
    )
    add_screenshot_placeholder(doc, "public/images/roles folder open in File Explorer")
    add_screenshot_placeholder(doc, "src/data/roles.ts with the image field highlighted")

    add_heading(doc, "3.2  Replace the favicon / logo", level=2)
    add_paragraph(doc, "The site's favicon is public/axis.svg. Replace the file (keep the same name) and the browser tab icon updates automatically.")
    add_code_block(
        doc,
        '<!-- index.html line 5 -->\n'
        '<link rel="icon" type="image/svg+xml" href="/axis.svg" />\n',
    )

    add_heading(doc, "3.3  Add a brand-new image inside a page", level=2)
    add_paragraph(doc, "Inside any component or page (.tsx file), use a normal <img> tag. The src must start with a forward slash because images come from /public.")
    add_code_block(
        doc,
        '<img\n'
        '  src="/images/your-new-image.png"\n'
        '  alt="Short description for accessibility"\n'
        '  className="w-full h-64 object-cover rounded-xl"\n'
        '/>\n',
    )
    add_note(doc, "Always provide a meaningful alt='' attribute — it helps SEO and screen readers.")

    doc.add_page_break()

    # ───────── 4. Text ─────────
    add_heading(doc, "4. How to Change Text", level=1)
    add_paragraph(
        doc,
        "There are two kinds of text in the site: (a) text that lives directly inside "
        "a page/component file, and (b) text driven by a data file like src/data/roles.ts.",
    )

    add_heading(doc, "4.1  Edit text in a page (e.g. the hero headline)", level=2)
    add_paragraph(doc, "File: src/pages/LandingPage.tsx")
    add_paragraph(doc, "Look for the text inside JSX tags such as <h1>, <h2>, <p>, etc.")
    add_code_block(
        doc,
        '<h1 className="text-white text-4xl sm:text-5xl md:text-7xl font-bold ...">\n'
        '  Hire the Top 1% of Global Talent and Save 80% on Payroll\n'
        '</h1>\n'
        '<p className="text-white/60 ...">\n'
        '  We\'ve saved 4,000+ companies over $300 million in payroll …\n'
        '</p>\n',
    )
    add_step(doc, 1, "Open the page file under src/pages/.")
    add_step(doc, 2, "Find the headline or paragraph inside the JSX.")
    add_step(doc, 3, "Type your new wording between the opening and closing tags.")
    add_step(doc, 4, "Save the file — the dev server hot-reloads it instantly.")
    add_screenshot_placeholder(doc, "LandingPage.tsx with the <h1> headline selected")

    add_heading(doc, "4.2  Edit text in a data file (role cards, steps, benefits)", level=2)
    add_paragraph(doc, "Many lists are kept in arrays so you don't have to dig through JSX. Examples:")
    add_code_block(
        doc,
        "// src/data/roles.ts\n"
        "{\n"
        "  title: 'Executive Assistant',          // ← edit title here\n"
        "  desc:  'Hire trusted, experienced …',  // ← edit description here\n"
        "  savings: '86%',                        // ← edit savings badge here\n"
        "  image: '/images/roles/executive-assistant.png'\n"
        "},\n",
    )
    add_code_block(
        doc,
        "// src/pages/LandingPage.tsx (top of file)\n"
        "const benefits = [\n"
        "  'College-educated professionals',\n"
        "  'Fluent English speakers',\n"
        "  'Expertise across tech, finance, marketing & more',\n"
        "  // ↑ Add, remove, or rewrite any line here\n"
        "]\n",
    )
    add_note(doc, "Keep quotes balanced. If your text contains an apostrophe, either escape it ( \\' ) or use double quotes around the string.")

    doc.add_page_break()

    # ───────── 5. Width ─────────
    add_heading(doc, "5. How to Change Width (Sizing & Layout)", level=1)
    add_paragraph(
        doc,
        "Width is controlled by Tailwind utility classes inside the className=\"…\" "
        "attribute of an element. Below are the classes you'll meet most often.",
    )

    add_heading(doc, "5.1  Common width classes", level=2)
    table = doc.add_table(rows=7, cols=2)
    table.style = "Light Grid Accent 1"
    rows = [
        ("Class", "Meaning"),
        ("w-full", "100% of the parent's width"),
        ("w-1/2, w-1/3, w-2/3", "50%, 33%, 66%"),
        ("w-64, w-96", "16rem (256px), 24rem (384px)"),
        ("max-w-2xl, max-w-4xl, max-w-7xl", "Caps the width at 42rem / 56rem / 80rem"),
        ("mx-auto", "Centers the element horizontally"),
        ("px-6, px-12", "Horizontal padding (1.5rem / 3rem)"),
    ]
    for r, (a, b) in enumerate(rows):
        table.rows[r].cells[0].text = a
        table.rows[r].cells[1].text = b
        if r == 0:
            for c in table.rows[r].cells:
                shade_cell(c, "DCE6F1")
                for run in c.paragraphs[0].runs:
                    run.bold = True

    add_heading(doc, "5.2  Example — making a section narrower", level=2)
    add_paragraph(doc, "Before:")
    add_code_block(
        doc,
        '<div className="max-w-7xl mx-auto px-6">\n'
        '  …content…\n'
        '</div>\n',
    )
    add_paragraph(doc, "After (tighter content area):")
    add_code_block(
        doc,
        '<div className="max-w-4xl mx-auto px-6">\n'
        '  …content…\n'
        '</div>\n',
    )
    add_screenshot_placeholder(doc, "Browser preview: before vs after width change")

    add_heading(doc, "5.3  Responsive widths", level=2)
    add_paragraph(doc, "Tailwind lets you change a class only at certain screen sizes by adding a prefix:")
    add_code_block(
        doc,
        '<!-- Full width on mobile, half-width on tablets and larger -->\n'
        '<div className="w-full md:w-1/2"> … </div>\n'
        '\n'
        '<!-- Prefix list:\n'
        '   sm: ≥640px   md: ≥768px   lg: ≥1024px   xl: ≥1280px   2xl: ≥1536px -->\n',
    )

    add_heading(doc, "5.4  Setting a custom pixel width", level=2)
    add_paragraph(doc, "If no Tailwind class fits, use an inline style:")
    add_code_block(
        doc,
        '<img src="/images/logo.png" alt="Logo" style={{ width: 220 }} />\n',
    )

    doc.add_page_break()

    # ───────── 6. Cheat sheet ─────────
    add_heading(doc, "6. Quick Reference Cheat-Sheet", level=1)
    table = doc.add_table(rows=6, cols=3)
    table.style = "Light Grid Accent 1"
    headers = ["What to change", "File to edit", "Where (line / key)"]
    rows = [
        ("Font (import)",        "index.html",            "Line 10 — <link href=…>"),
        ("Font (Tailwind class)","tailwind.config.js",    "Lines 36–39 — fontFamily"),
        ("Default body font",    "src/index.css",         "Line 18 — body { font-family }"),
        ("Role card image/text", "src/data/roles.ts",     "image / title / desc fields"),
        ("Page headline / copy", "src/pages/<Page>.tsx",  "Inside <h1>, <h2>, <p> tags"),
    ]
    for c, h in enumerate(headers):
        cell = table.rows[0].cells[c]
        cell.text = h
        shade_cell(cell, "DCE6F1")
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for r, row in enumerate(rows, 1):
        for c, val in enumerate(row):
            table.rows[r].cells[c].text = val

    doc.add_paragraph()
    add_heading(doc, "Workflow tips", level=2)
    for tip in [
        "Run `npm run dev` to start the local preview server. Changes appear live.",
        "Save the file (Ctrl+S) to trigger hot-reload in the browser.",
        "If something looks broken, check the browser DevTools console (F12) for red errors.",
        "Always commit working changes with git before doing big rewrites.",
    ]:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(tip)

    out_path = "Axis_Website_User_Manual.docx"
    doc.save(out_path)
    print(f"WROTE: {out_path}")


if __name__ == "__main__":
    main()
